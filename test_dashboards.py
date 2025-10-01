"""
Test script to validate dashboard functionality
"""

import os
import openpyxl
from dashboard_generator import DashboardGenerator
from sample_data import SampleDataGenerator


def test_dashboard_generation():
    """Test that all dashboards are generated correctly"""
    print("Testing dashboard generation...")
    
    # Clean up existing test output
    test_dir = '/tmp/test_dashboards'
    if os.path.exists(test_dir):
        import shutil
        shutil.rmtree(test_dir)
    
    generator = DashboardGenerator(output_dir=test_dir)
    
    # Test issue tracker generation
    issue_tracker_path = generator.create_issue_tracker()
    assert os.path.exists(issue_tracker_path), "Issue tracker not created"
    print("✓ Issue tracker created successfully")
    
    # Test client launch tracker generation
    launch_tracker_path = generator.create_client_launch_tracker()
    assert os.path.exists(launch_tracker_path), "Client launch tracker not created"
    print("✓ Client launch tracker created successfully")
    
    # Test integrated dashboard generation
    integrated_path = generator.create_integrated_dashboard()
    assert os.path.exists(integrated_path), "Integrated dashboard not created"
    print("✓ Integrated dashboard created successfully")
    
    return True


def test_issue_tracker_structure():
    """Test issue tracker structure"""
    print("\nTesting issue tracker structure...")
    
    test_dir = '/tmp/test_dashboards'
    filepath = os.path.join(test_dir, 'issue_tracker_dashboard.xlsx')
    
    wb = openpyxl.load_workbook(filepath)
    
    # Check sheets
    expected_sheets = ['Issue Tracker', 'Issue Summary', 'KPI Dashboard']
    for sheet in expected_sheets:
        assert sheet in wb.sheetnames, f"Sheet '{sheet}' not found"
    print(f"✓ All expected sheets present: {expected_sheets}")
    
    # Check Issue Tracker columns
    ws = wb['Issue Tracker']
    expected_columns = [
        'Issue ID', 'Title', 'Description', 'Priority', 'Status',
        'Assigned To', 'Reporter', 'Created Date', 'Due Date',
        'Resolution Date', 'Category', 'Tags'
    ]
    
    actual_columns = [ws.cell(1, col).value for col in range(1, len(expected_columns) + 1)]
    assert actual_columns == expected_columns, "Issue Tracker columns don't match"
    print(f"✓ Issue Tracker has correct columns ({len(expected_columns)} columns)")
    
    # Check that there's sample data
    assert ws.max_row > 1, "No sample data in Issue Tracker"
    print(f"✓ Issue Tracker has sample data ({ws.max_row - 1} rows)")
    
    return True


def test_launch_tracker_structure():
    """Test launch tracker structure"""
    print("\nTesting launch tracker structure...")
    
    test_dir = '/tmp/test_dashboards'
    filepath = os.path.join(test_dir, 'client_launch_tracker.xlsx')
    
    wb = openpyxl.load_workbook(filepath)
    
    # Check sheets
    expected_sheets = ['Launch Tracker', 'Timeline']
    for sheet in expected_sheets:
        assert sheet in wb.sheetnames, f"Sheet '{sheet}' not found"
    print(f"✓ All expected sheets present: {expected_sheets}")
    
    # Check Launch Tracker columns
    ws = wb['Launch Tracker']
    expected_columns = [
        'Client Name', 'Project Name', 'Launch Date', 'Status',
        'Project Manager', 'Budget', 'Completion %', 'Phase',
        'Risk Level', 'Notes'
    ]
    
    actual_columns = [ws.cell(1, col).value for col in range(1, len(expected_columns) + 1)]
    assert actual_columns == expected_columns, "Launch Tracker columns don't match"
    print(f"✓ Launch Tracker has correct columns ({len(expected_columns)} columns)")
    
    # Check that there's sample data
    assert ws.max_row > 1, "No sample data in Launch Tracker"
    print(f"✓ Launch Tracker has sample data ({ws.max_row - 1} rows)")
    
    return True


def test_integrated_dashboard_structure():
    """Test integrated dashboard structure"""
    print("\nTesting integrated dashboard structure...")
    
    test_dir = '/tmp/test_dashboards'
    filepath = os.path.join(test_dir, 'integrated_dashboard.xlsx')
    
    wb = openpyxl.load_workbook(filepath)
    
    # Check sheets
    expected_sheets = ['Dashboard Overview', 'Issue Tracker', 'Launch Tracker', 'KPI Dashboard']
    for sheet in expected_sheets:
        assert sheet in wb.sheetnames, f"Sheet '{sheet}' not found"
    print(f"✓ All expected sheets present: {expected_sheets}")
    
    # Check Dashboard Overview has content
    ws = wb['Dashboard Overview']
    assert ws['A1'].value is not None, "Dashboard Overview is empty"
    print("✓ Dashboard Overview has content")
    
    return True


def test_sample_data_generator():
    """Test sample data generator"""
    print("\nTesting sample data generator...")
    
    generator = SampleDataGenerator()
    
    # Test issue generation
    issues = generator.generate_issues(5)
    assert len(issues) == 5, "Wrong number of issues generated"
    assert all('issue_id' in issue for issue in issues), "Issues missing required fields"
    print(f"✓ Generated {len(issues)} sample issues")
    
    # Test client launch generation
    launches = generator.generate_client_launches(3)
    assert len(launches) == 3, "Wrong number of launches generated"
    assert all('client_name' in launch for launch in launches), "Launches missing required fields"
    print(f"✓ Generated {len(launches)} sample launches")
    
    # Test KPI data generation
    kpis = generator.generate_kpi_data()
    assert 'total_issues' in kpis, "KPI data missing required fields"
    assert 'weekly_trends' in kpis, "KPI data missing weekly trends"
    print("✓ Generated sample KPI data")
    
    return True


def test_color_formatting():
    """Test that color formatting is applied"""
    print("\nTesting color formatting...")
    
    test_dir = '/tmp/test_dashboards'
    filepath = os.path.join(test_dir, 'integrated_dashboard.xlsx')
    
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Issue Tracker']
    
    # Check that headers have background color
    header_cell = ws.cell(1, 1)
    assert header_cell.fill.start_color.rgb is not None, "Header cell has no background color"
    print("✓ Headers have color formatting")
    
    # Check that priority cells have conditional formatting
    # Find a row with priority data
    priority_col = 4
    priority_found = False
    for row in range(2, min(ws.max_row + 1, 10)):
        cell = ws.cell(row, priority_col)
        if cell.value in ['Critical', 'High', 'Medium', 'Low']:
            if cell.fill.start_color.rgb:
                priority_found = True
                break
    
    assert priority_found, "Priority cells don't have color formatting"
    print("✓ Priority cells have color formatting")
    
    return True


def run_all_tests():
    """Run all tests"""
    print("=" * 60)
    print("FAST Project Dashboard - Test Suite")
    print("=" * 60)
    
    tests = [
        ("Dashboard Generation", test_dashboard_generation),
        ("Issue Tracker Structure", test_issue_tracker_structure),
        ("Launch Tracker Structure", test_launch_tracker_structure),
        ("Integrated Dashboard Structure", test_integrated_dashboard_structure),
        ("Sample Data Generator", test_sample_data_generator),
        ("Color Formatting", test_color_formatting),
    ]
    
    passed = 0
    failed = 0
    
    for test_name, test_func in tests:
        try:
            if test_func():
                passed += 1
        except Exception as e:
            print(f"✗ {test_name} failed: {str(e)}")
            failed += 1
    
    print("\n" + "=" * 60)
    print(f"Test Results: {passed} passed, {failed} failed")
    print("=" * 60)
    
    if failed == 0:
        print("\n✓ All tests passed!")
        return True
    else:
        print(f"\n✗ {failed} test(s) failed")
        return False


if __name__ == '__main__':
    success = run_all_tests()
    exit(0 if success else 1)
