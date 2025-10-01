"""
Example usage of the dashboard system
Demonstrates how to programmatically add data to existing dashboards
"""

import openpyxl
from openpyxl.styles import PatternFill, Font
from datetime import datetime
import os


def add_issue_to_tracker(filepath, issue_data):
    """
    Add a new issue to an existing issue tracker
    
    Args:
        filepath: Path to the Excel file
        issue_data: Dictionary containing issue information
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Issue Tracker']
    
    # Find the next empty row
    next_row = ws.max_row + 1
    
    # Add data
    columns = ['issue_id', 'title', 'description', 'priority', 'status',
               'assigned_to', 'reporter', 'created_date', 'due_date',
               'resolution_date', 'category', 'tags']
    
    for col_idx, col_name in enumerate(columns, 1):
        value = issue_data.get(col_name, '')
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        
        # Apply priority formatting
        if col_idx == 4:  # Priority column
            if value == 'Critical':
                cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                cell.font = Font(color='FFFFFF', bold=True)
            elif value == 'High':
                cell.fill = PatternFill(start_color='FFA500', end_color='FFA500', fill_type='solid')
        
        # Apply status formatting
        if col_idx == 5:  # Status column
            if value == 'Resolved':
                cell.fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')
            elif value == 'In Progress':
                cell.fill = PatternFill(start_color='87CEEB', end_color='87CEEB', fill_type='solid')
    
    wb.save(filepath)
    print(f"Added issue {issue_data.get('issue_id', 'N/A')} to {filepath}")


def add_client_launch(filepath, launch_data):
    """
    Add a new client launch to the tracker
    
    Args:
        filepath: Path to the Excel file
        launch_data: Dictionary containing launch information
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Launch Tracker']
    
    # Find the next empty row
    next_row = ws.max_row + 1
    
    # Add data
    columns = ['client_name', 'project_name', 'launch_date', 'status',
               'project_manager', 'budget', 'completion_pct', 'phase',
               'risk_level', 'notes']
    
    for col_idx, col_name in enumerate(columns, 1):
        value = launch_data.get(col_name, '')
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        
        # Apply status formatting
        if col_idx == 4:  # Status column
            if value == 'Completed':
                cell.fill = PatternFill(start_color='00B050', end_color='00B050', fill_type='solid')
                cell.font = Font(color='FFFFFF', bold=True)
            elif value == 'On Track':
                cell.fill = PatternFill(start_color='92D050', end_color='92D050', fill_type='solid')
            elif value == 'At Risk':
                cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                cell.font = Font(color='FFFFFF', bold=True)
    
    wb.save(filepath)
    print(f"Added client launch for {launch_data.get('client_name', 'N/A')} to {filepath}")


def read_issues_from_tracker(filepath, status_filter=None):
    """
    Read issues from the tracker with optional filtering
    
    Args:
        filepath: Path to the Excel file
        status_filter: Optional status to filter by (e.g., 'Open', 'In Progress')
    
    Returns:
        List of issue dictionaries
    """
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Issue Tracker']
    
    # Get headers
    headers = [cell.value for cell in ws[1]]
    
    # Read data
    issues = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        issue = dict(zip(headers, row))
        
        # Apply filter if specified
        if status_filter is None or issue.get('Status') == status_filter:
            issues.append(issue)
    
    return issues


def example_usage():
    """Example of using the dashboard system"""
    
    print("=" * 60)
    print("FAST Project Dashboard - Example Usage")
    print("=" * 60)
    
    # Check if dashboards exist
    integrated_path = 'output/integrated_dashboard.xlsx'
    if not os.path.exists(integrated_path):
        print("\nError: Dashboards not found!")
        print("Please run 'python dashboard_generator.py' first.")
        return
    
    # Example 1: Add a new issue
    print("\n1. Adding a new issue...")
    new_issue = {
        'issue_id': 'ISS-006',
        'title': 'Email Notification System',
        'description': 'Implement email notifications for system alerts',
        'priority': 'High',
        'status': 'Open',
        'assigned_to': 'Developer Team',
        'reporter': 'Product Manager',
        'created_date': datetime.now().strftime('%Y-%m-%d'),
        'due_date': '2024-02-15',
        'resolution_date': '',
        'category': 'Feature',
        'tags': 'email,notification,feature'
    }
    
    add_issue_to_tracker(integrated_path, new_issue)
    
    # Example 2: Add a new client launch
    print("\n2. Adding a new client launch...")
    new_launch = {
        'client_name': 'NewTech Solutions',
        'project_name': 'E-commerce Platform',
        'launch_date': '2024-06-30',
        'status': 'Planning',
        'project_manager': 'Project Lead',
        'budget': '$175,000',
        'completion_pct': 10,
        'phase': 'Planning',
        'risk_level': 'Medium',
        'notes': 'Initial planning phase started'
    }
    
    add_client_launch(integrated_path, new_launch)
    
    # Example 3: Read and display open issues
    print("\n3. Reading open issues...")
    open_issues = read_issues_from_tracker(integrated_path, status_filter='Open')
    
    print(f"\nFound {len(open_issues)} open issue(s):")
    for issue in open_issues:
        print(f"  - {issue['Issue ID']}: {issue['Title']} (Priority: {issue['Priority']})")
    
    print("\n" + "=" * 60)
    print("Example completed successfully!")
    print("=" * 60)
    print(f"\nCheck the updated file: {integrated_path}")


if __name__ == '__main__':
    example_usage()
