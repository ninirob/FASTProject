"""
Display information about the generated dashboards
"""

import openpyxl
import os
from datetime import datetime


def display_dashboard_info(filepath):
    """Display information about a dashboard file"""
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return
    
    wb = openpyxl.load_workbook(filepath)
    filename = os.path.basename(filepath)
    
    print(f"\n{'='*70}")
    print(f"Dashboard: {filename}")
    print(f"{'='*70}")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        print(f"\n📊 Sheet: {sheet_name}")
        print(f"   Rows: {ws.max_row}")
        print(f"   Columns: {ws.max_column}")
        
        # Display headers if available
        if ws.max_row > 0:
            headers = []
            for col in range(1, min(ws.max_column + 1, 11)):
                header = ws.cell(1, col).value
                if header:
                    headers.append(str(header))
            
            if headers:
                print(f"   Headers: {', '.join(headers[:5])}{'...' if len(headers) > 5 else ''}")
        
        # Count data rows (excluding header)
        data_rows = ws.max_row - 1 if ws.max_row > 1 else 0
        if data_rows > 0:
            print(f"   Data rows: {data_rows}")


def display_statistics():
    """Display statistics about all dashboards"""
    output_dir = 'output'
    
    if not os.path.exists(output_dir):
        print(f"Output directory not found: {output_dir}")
        print("Please run: python dashboard_generator.py")
        return
    
    print("=" * 70)
    print("FAST Project Dashboard System - Overview")
    print("=" * 70)
    print(f"\nGenerated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    # List all Excel files
    excel_files = [f for f in os.listdir(output_dir) if f.endswith('.xlsx')]
    
    print(f"\nTotal dashboards: {len(excel_files)}")
    
    for filename in sorted(excel_files):
        filepath = os.path.join(output_dir, filename)
        display_dashboard_info(filepath)
    
    print("\n" + "=" * 70)
    print("Dashboard Capabilities:")
    print("=" * 70)
    
    capabilities = [
        ("Issue Tracking", "Track bugs, tasks, and issues with priority and status"),
        ("Client Management", "Monitor client projects and launch timelines"),
        ("KPI Monitoring", "View key metrics and performance indicators"),
        ("Visual Analytics", "Charts and graphs for trend analysis"),
        ("Status Tracking", "Color-coded status and priority indicators"),
        ("Timeline Management", "Track milestones and deadlines"),
        ("Budget Tracking", "Monitor project budgets and completion"),
        ("Risk Assessment", "Identify and track project risks"),
    ]
    
    for capability, description in capabilities:
        print(f"\n✓ {capability}")
        print(f"  {description}")
    
    print("\n" + "=" * 70)
    print("Quick Actions:")
    print("=" * 70)
    print("\n1. Generate fresh dashboards:")
    print("   python dashboard_generator.py")
    print("\n2. Add data programmatically:")
    print("   python example_usage.py")
    print("\n3. Test with sample data:")
    print("   python sample_data.py")
    print("\n4. Run tests:")
    print("   python test_dashboards.py")
    print("\n" + "=" * 70)


def main():
    display_statistics()


if __name__ == '__main__':
    main()
