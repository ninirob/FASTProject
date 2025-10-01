"""
Demo script showing the complete dashboard system in action
"""

import os
from dashboard_generator import DashboardGenerator
from sample_data import SampleDataGenerator
import openpyxl


def run_demo():
    """Run a complete demonstration of the dashboard system"""
    
    print("=" * 70)
    print("FAST Project Dashboard System - Complete Demo")
    print("=" * 70)
    
    # Step 1: Generate dashboards
    print("\n[Step 1] Generating dashboards...")
    print("-" * 70)
    
    generator = DashboardGenerator()
    
    print("\n📊 Creating Issue Tracker Dashboard...")
    issue_path = generator.create_issue_tracker()
    
    print("\n📊 Creating Client Launch Tracker...")
    launch_path = generator.create_client_launch_tracker()
    
    print("\n📊 Creating Integrated Dashboard...")
    integrated_path = generator.create_integrated_dashboard()
    
    # Step 2: Show dashboard statistics
    print("\n[Step 2] Dashboard Statistics")
    print("-" * 70)
    
    dashboards = [
        ('Issue Tracker Dashboard', issue_path),
        ('Client Launch Tracker', launch_path),
        ('Integrated Dashboard', integrated_path),
    ]
    
    for name, path in dashboards:
        wb = openpyxl.load_workbook(path)
        print(f"\n{name}:")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            data_rows = ws.max_row - 1 if ws.max_row > 1 else 0
            print(f"  • {sheet_name}: {data_rows} rows of data")
    
    # Step 3: Demonstrate data connections
    print("\n[Step 3] Demonstrating Dashboard Connections")
    print("-" * 70)
    
    print("\n📋 Reading Issue Tracker Data...")
    wb = openpyxl.load_workbook(integrated_path)
    ws = wb['Issue Tracker']
    
    # Count issues by status
    status_counts = {}
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row, 5).value
        if status:
            status_counts[status] = status_counts.get(status, 0) + 1
    
    print("\nIssue Status Distribution:")
    for status, count in status_counts.items():
        print(f"  • {status}: {count}")
    
    # Count issues by priority
    priority_counts = {}
    for row in range(2, ws.max_row + 1):
        priority = ws.cell(row, 4).value
        if priority:
            priority_counts[priority] = priority_counts.get(priority, 0) + 1
    
    print("\nIssue Priority Distribution:")
    for priority, count in sorted(priority_counts.items(), key=lambda x: x[1], reverse=True):
        print(f"  • {priority}: {count}")
    
    print("\n📋 Reading Client Launch Data...")
    ws = wb['Launch Tracker']
    
    # Count projects by status
    project_status_counts = {}
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row, 4).value
        if status:
            project_status_counts[status] = project_status_counts.get(status, 0) + 1
    
    print("\nProject Status Distribution:")
    for status, count in project_status_counts.items():
        print(f"  • {status}: {count}")
    
    # Calculate average completion
    completion_sum = 0
    completion_count = 0
    for row in range(2, ws.max_row + 1):
        completion = ws.cell(row, 7).value
        if isinstance(completion, (int, float)):
            completion_sum += completion
            completion_count += 1
    
    if completion_count > 0:
        avg_completion = completion_sum / completion_count
        print(f"\nAverage Project Completion: {avg_completion:.1f}%")
    
    # Step 4: Show KPI Dashboard
    print("\n[Step 4] KPI Dashboard Summary")
    print("-" * 70)
    
    ws = wb['KPI Dashboard']
    
    print("\nKey Performance Indicators:")
    kpi_cells = [
        ('B4', 'Total Issues'),
        ('D4', 'Open Issues'),
        ('F4', 'Resolved Issues'),
        ('B7', 'Avg Resolution Time'),
        ('D7', 'Critical Issues'),
        ('F7', 'On-Time Delivery Rate'),
    ]
    
    for cell_ref, label in kpi_cells:
        value = ws[cell_ref].value
        print(f"  • {label}: {value}")
    
    # Step 5: Show integration capabilities
    print("\n[Step 5] Integration Capabilities")
    print("-" * 70)
    
    print("\n✓ All dashboards connect seamlessly:")
    print("  • Issue Tracker feeds into KPI Dashboard")
    print("  • Client Launch Tracker provides project metrics")
    print("  • Integrated Dashboard combines all views")
    print("  • Charts update automatically with data changes")
    print("  • Color coding consistent across all dashboards")
    print("  • Excel formulas can link between sheets")
    
    # Step 6: Usage examples
    print("\n[Step 6] Next Steps & Usage")
    print("-" * 70)
    
    print("\n1. Open the dashboards in Excel:")
    print(f"   • {os.path.basename(integrated_path)}")
    
    print("\n2. Customize the data:")
    print("   • Add your own issues and projects")
    print("   • Update statuses and priorities")
    print("   • Track progress over time")
    
    print("\n3. Use programmatically:")
    print("   • Run: python example_usage.py")
    print("   • Integrate with your own systems")
    print("   • Automate report generation")
    
    print("\n4. Generate reports:")
    print("   • Export specific sheets as PDF")
    print("   • Share with stakeholders")
    print("   • Archive for historical tracking")
    
    # Summary
    print("\n" + "=" * 70)
    print("Demo Complete!")
    print("=" * 70)
    
    total_files = len(dashboards)
    total_sheets = sum(len(openpyxl.load_workbook(path).sheetnames) for _, path in dashboards)
    
    print(f"\n✓ Generated {total_files} dashboard files")
    print(f"✓ Created {total_sheets} interconnected sheets")
    print(f"✓ Sample data populated and ready to use")
    print(f"✓ All dashboards connected and integrated")
    
    print("\n📁 Files created in: output/")
    print("📖 See README.md for detailed documentation")
    print("🚀 See QUICKSTART.md for getting started guide")
    
    print("\n" + "=" * 70)


def main():
    """Main entry point"""
    try:
        run_demo()
    except Exception as e:
        print(f"\n❌ Error during demo: {str(e)}")
        raise


if __name__ == '__main__':
    main()
